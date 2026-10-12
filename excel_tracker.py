import os
from datetime import datetime
from typing import List, Optional
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.config import APPLICATIONS_FILE
from backend.models.schemas import ApplicationRecord

COLUMNS = [
    "Application ID", "Date", "Company", "Role", "Job URL",
    "Fit Score", "Gate Status", "ATS Score", "HR Score",
    "Evidence Match %", "Resume DOCX", "Cover Letter DOCX",
    "Status", "Notes"
]

HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ZEBRA_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

def init_excel_tracker():
    if not APPLICATIONS_FILE.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Applications"
        ws.views.sheetView[0].showGridLines = True
        
        # Write headers
        ws.append(COLUMNS)
        for col_num in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER
        ws.row_dimensions[1].height = 28
        wb.save(str(APPLICATIONS_FILE))

def auto_format_sheet(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(12, min(max_len + 3, 40))

def add_application(record: ApplicationRecord):
    init_excel_tracker()
    wb = load_workbook(str(APPLICATIONS_FILE))
    ws = wb["Applications"]
    
    # Check if record id already exists; update if so
    found_row = None
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == record.id:
            found_row = row
            break
            
    row_data = [
        record.id,
        record.date,
        record.company,
        record.role,
        record.url,
        record.fit_score,
        record.gate_status,
        record.ats_score,
        record.hr_score,
        f"{record.evidence_match}%",
        record.resume_file,
        record.cover_letter_file,
        record.status,
        record.notes
    ]
    
    target_row = found_row if found_row else ws.max_row + 1
    
    for col_idx, val in enumerate(row_data, 1):
        cell = ws.cell(row=target_row, column=col_idx, value=val)
        cell.font = Font(name="Calibri", size=10)
        cell.border = THIN_BORDER
        cell.fill = ZEBRA_FILL if (target_row % 2 == 0) else WHITE_FILL
        if col_idx in [1, 2, 6, 7, 8, 9, 10, 13]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            
    ws.row_dimensions[target_row].height = 22
    auto_format_sheet(ws)
    wb.save(str(APPLICATIONS_FILE))

def get_all_applications() -> List[ApplicationRecord]:
    init_excel_tracker()
    records = []
    wb = load_workbook(str(APPLICATIONS_FILE))
    ws = wb["Applications"]
    
    for row in range(2, ws.max_row + 1):
        app_id = ws.cell(row=row, column=1).value
        if not app_id:
            continue
        date_val = str(ws.cell(row=row, column=2).value or "")
        company = str(ws.cell(row=row, column=3).value or "")
        role = str(ws.cell(row=row, column=4).value or "")
        url = str(ws.cell(row=row, column=5).value or "")
        fit_score = int(ws.cell(row=row, column=6).value or 0)
        gate_status = str(ws.cell(row=row, column=7).value or "Passed")
        ats_score = int(ws.cell(row=row, column=8).value or 0)
        hr_score = int(ws.cell(row=row, column=9).value or 0)
        
        ev_str = str(ws.cell(row=row, column=10).value or "0").replace("%", "").strip()
        try:
            evidence_match = float(ev_str)
        except ValueError:
            evidence_match = 0.0
            
        resume_file = str(ws.cell(row=row, column=11).value or "")
        cover_letter_file = str(ws.cell(row=row, column=12).value or "")
        status = str(ws.cell(row=row, column=13).value or "Ready")
        notes = str(ws.cell(row=row, column=14).value or "")
        
        records.append(ApplicationRecord(
            id=str(app_id),
            date=date_val,
            company=company,
            role=role,
            url=url,
            fit_score=fit_score,
            gate_status=gate_status,
            ats_score=ats_score,
            hr_score=hr_score,
            evidence_match=evidence_match,
            resume_file=resume_file,
            cover_letter_file=cover_letter_file,
            status=status,
            notes=notes
        ))
    return records

def update_application_status(app_id: str, new_status: str, notes: Optional[str] = None):
    init_excel_tracker()
    wb = load_workbook(str(APPLICATIONS_FILE))
    ws = wb["Applications"]
    
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row=row, column=1).value) == str(app_id):
            ws.cell(row=row, column=13, value=new_status)
            if notes is not None:
                ws.cell(row=row, column=14, value=notes)
            wb.save(str(APPLICATIONS_FILE))
            return True
    return False
